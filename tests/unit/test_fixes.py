"""Tests for the entrepreneur-stage fixes: findings (X1), data checks (X2), workbook (X3)."""
import io as _io
from datetime import datetime

import openpyxl
import pandas as pd
import pytest

from smartlytics.core import checks, findings, sales, workbook
from smartlytics.core.audit import AuditEntry


# ---- X2 data checks -----------------------------------------------------------
def _shop() -> pd.DataFrame:
    return pd.DataFrame({
        "date": pd.to_datetime(["2026-01-01", "2026-01-02", "2099-01-01", "2026-01-03"]),
        "quantity": [2, 1, 3, -1],
        "unit_price": [10.0, 20.0, 5.0, 10.0],
        "unit_cost": [4.0, 25.0, 2.0, None],
        "gross_sales": [20.0, 200.0, 15.0, -10.0],    # row 1: typo (should be 20)
        "discount_pct": [0, 10, 150, 0],              # row 2: impossible discount
        "revenue": [20.0, 18.0, -7.5, -10.0],
    })


def _mapping() -> sales.SalesMapping:
    return sales.SalesMapping(revenue="revenue", cost="unit_cost", quantity="quantity", date="date")


def test_extra_columns_found():
    assert checks.find_extra_columns(_shop(), _mapping()) == {
        "price": "unit_price", "gross": "gross_sales", "discount": "discount_pct"}


def test_each_check_flags_the_planted_row():
    found = {c.name: c for c in checks.run_checks(_shop(), _mapping(), today=datetime(2026, 10, 22))}
    assert found["gross_sales = quantity × unit_price"].rows.index.tolist() == [1]
    assert found["discount_pct between 0 and 100"].rows.index.tolist() == [2]
    assert found["unit_price above unit_cost"].rows.index.tolist() == [1]
    assert found["unit_cost present"].flagged == 1
    assert found["quantity positive"].rows.index.tolist() == [3]
    assert found["date not in the future"].rows.index.tolist() == [2]


def test_revenue_check_skips_rows_with_wrong_gross():
    found = {c.name: c for c in checks.run_checks(_shop(), _mapping())}
    rev = found["revenue = gross_sales after discount_pct"]
    assert 1 not in rev.rows.index          # gross already wrong there: not double-counted
    assert rev.checked == 3


def test_checks_on_sample_find_the_planted_typos(sample_clean):
    found = {c.name: c for c in checks.run_checks(sample_clean, sales.guess_mapping(sample_clean))}
    assert found["gross_sales = quantity × unit_price"].flagged == 25
    assert found["revenue = gross_sales after discount_pct"].flagged == 0
    assert found["quantity positive"].flagged == 12


def test_no_extra_columns_still_runs_basic_checks():
    df = pd.DataFrame({"revenue": [1.0, -2.0]})
    names = [c.name for c in checks.run_checks(df, sales.SalesMapping(revenue="revenue"))]
    assert names == ["revenue positive"]


# ---- X1 findings ----------------------------------------------------------------
def test_findings_are_computed_facts(sample_result):
    facts = findings.findings(sample_result)
    assert 4 <= len(facts) <= 6
    high_share = sample_result.df.loc[sample_result.df["tier"] == "High", "profit"].sum() / \
        sample_result.df.loc[sample_result.df["tier"].notna(), "profit"].sum() * 100
    assert f"{high_share:.0f}%" in facts[0]
    assert any("Laptop" in f for f in facts)
    assert any("top 10 customers" in f for f in facts)


def test_findings_on_revenue_only():
    df = pd.DataFrame({"revenue": [10.0, 20.0, 30.0, 40.0], "product": ["a", "b", "a", "b"]})
    r = sales.analyse(df, sales.SalesMapping(revenue="revenue", product="product"))
    facts = findings.findings(r)
    assert all("profit" not in f for f in facts)
    assert any("b earns" in f for f in facts)


def test_customer_concentration_wording():
    df = pd.DataFrame({"revenue": [1.0] * 500, "customer": [f"c{i}" for i in range(500)]})
    r = sales.analyse(df, sales.SalesMapping(revenue="revenue", customer="customer"))
    assert any(f.startswith("No customer dominates") for f in findings.findings(r))


# ---- X3 workbook ----------------------------------------------------------------
def test_workbook_sheets_and_numbers(sample_clean, sample_result):
    data = workbook.build_workbook("Sample", sample_clean, [AuditEntry("x", "Removed 40 duplicate rows", 5040, 5000)],
                                   sample_result)
    wb = openpyxl.load_workbook(_io.BytesIO(data))
    assert wb.sheetnames == ["Summary", "Tiers", "Tier boundaries", "Top performers", "Data checks",
                             "Rows to look at", "Cleaning steps", "Data"]
    summary = {r[0]: r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True)}
    assert summary["Rows"] == "5,000"
    assert summary["Profit"] == f"{sample_result.totals['profit']:,.2f}"
    assert wb["Data"].max_row == 5001
    assert wb["Summary"]["A1"].font.bold


def test_workbook_without_sales(sample_clean):
    wb = openpyxl.load_workbook(_io.BytesIO(workbook.build_workbook("S", sample_clean, [], None)))
    assert wb.sheetnames == ["Summary", "Data"]


def test_workbook_skips_data_sheet_when_large(monkeypatch, sample_clean, sample_result):
    monkeypatch.setattr(workbook, "DATA_SHEET_MAX", 100)
    wb = openpyxl.load_workbook(_io.BytesIO(workbook.build_workbook("S", sample_clean, [], sample_result)))
    assert "Data" not in wb.sheetnames
    notes = [r[1] for r in wb["Summary"].iter_rows(min_row=2, values_only=True) if r[0] == "Note"]
    assert any("CSV download" in n for n in notes)


@pytest.mark.parametrize("tz", ["UTC", None])
def test_workbook_handles_timezone_dates(tz):
    df = pd.DataFrame({"d": pd.date_range("2026-01-01", periods=3, tz=tz), "revenue": [1.0, 2.0, 3.0]})
    assert workbook.build_workbook("S", df, [], None)
