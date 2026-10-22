"""Excel workbook export: the answers, in the tool people already use."""
from __future__ import annotations

import io

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from truesum.core import checks as checks_mod
from truesum.core import findings as findings_mod
from truesum.core import sales
from truesum.core.audit import AuditEntry

DATA_SHEET_MAX = 25_000   # writing cells is ~5k rows/s; bigger data goes in the CSV download
FLAGGED_MAX = 1_000
TOP_MAX = 5_000
HEADER_FILL = PatternFill("solid", fgColor="0F2027")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _strip_tz(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if isinstance(out[c].dtype, pd.DatetimeTZDtype):
            out[c] = out[c].dt.tz_localize(None)
    return out


def _style(ws, df: pd.DataFrame) -> None:
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    for i, col in enumerate(df.columns, start=1):
        sample = [str(col)] + [str(v) for v in df[col].head(200).tolist()]
        ws.column_dimensions[get_column_letter(i)].width = min(max(len(s) for s in sample) + 2, 60)
        if pd.api.types.is_float_dtype(df[col]):
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i, max_row=min(ws.max_row, DATA_SHEET_MAX + 1)):
                cell.number_format = "#,##0.00"
        elif pd.api.types.is_datetime64_any_dtype(df[col]):
            for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i, max_row=min(ws.max_row, DATA_SHEET_MAX + 1)):
                cell.number_format = "yyyy-mm-dd"


def _sheet(writer: pd.ExcelWriter, name: str, df: pd.DataFrame) -> None:
    df = _strip_tz(df)
    df.to_excel(writer, sheet_name=name, index=False)
    _style(writer.sheets[name], df)


def build_workbook(dataset_name: str, df: pd.DataFrame, audit: list[AuditEntry],
                   result: sales.SalesResult | None) -> bytes:
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        lines = [("Dataset", dataset_name), ("Rows", f"{len(df):,}"), ("Columns", str(df.shape[1]))]
        if result is not None:
            t = result.totals
            lines += [("Revenue", f"{t['revenue']:,.2f}")]
            if result.metric == "profit":
                lines += [("Profit", f"{t['profit']:,.2f}"), ("Margin", f"{t['margin_pct']:.1f}%")]
            lines += [("Rated on", result.metric), ("Low / Medium boundary", f"{result.thresholds.q25:,.2f}"),
                      ("Medium / High boundary", f"{result.thresholds.q75:,.2f}"), ("", "")]
            lines += [("Finding", f) for f in findings_mod.findings(result)]
            lines += [("Note", wtext) for wtext in result.warnings]
        if df.shape[0] > DATA_SHEET_MAX:
            lines.append(("Note", f"The full {len(df):,} rows are in the CSV download; this workbook holds the results."))
        _sheet(w, "Summary", pd.DataFrame(lines, columns=["item", "value"]))

        if result is not None:
            _sheet(w, "Tiers", sales.tier_summary(result))
            _sheet(w, "Tier boundaries", sales.threshold_table(result.thresholds))
            top = sales.top_performers(result)
            _sheet(w, "Top performers", top.head(TOP_MAX))
            if len(top) > TOP_MAX:
                lines_note = f"Top performers sheet shows the best {TOP_MAX:,} of {len(top):,} High-tier rows."
                w.sheets["Summary"].append(["Note", lines_note])
            found = checks_mod.run_checks(result.df, result.mapping)
            _sheet(w, "Data checks", checks_mod.summary(found))
            flagged = [c.rows.head(FLAGGED_MAX).assign(failed_check=c.name) for c in found if not c.passed]
            if flagged:
                _sheet(w, "Rows to look at", pd.concat(flagged, ignore_index=True))

        if audit:
            _sheet(w, "Cleaning steps", pd.DataFrame(
                [{"step": e.message, "rows before": e.rows_before, "rows after": e.rows_after,
                  "time (UTC)": e.at.strftime("%Y-%m-%d %H:%M")} for e in audit]))
        data = result.df if result is not None else df
        if len(data) <= DATA_SHEET_MAX:
            _sheet(w, "Data", data)
    return buf.getvalue()
